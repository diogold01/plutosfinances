"""
Motor independente de leitura e exportação de planilhas Excel `.xlsx` (`ExcelSpreadsheetHandler`).

Recursos:
- Carrega e preenche automaticamente os templates oficiais em `app/assets/excel_templates/`:
  - `template_despesas_excel.xlsx`
  - `template_receitas_excel.xlsx`
  - `template_transferencias_excel.xlsx`
  - `template_extrato_excel.xlsx`
  - `template_rel_group_cat_excel.xlsx`
- Permite gerar pasta de trabalho completa multi-abas (`Extrato`, `Despesas`, `Receitas`,
  `Transferencias`, `Categorias`) com estilização nativa OpenPyXL, formatação monetária
  `R$ #,##0.00` e ajuste automático de largura de colunas.
- Lê planilhas `.xlsx` de qualquer aba, detectando cabeçalhos e tipos automaticamente.
- Executa hooks externos (`ExternalSpreadsheetFormatterHook`) para integração com ferramentas
  externas de formatação de planilhas (Requisito 3).
"""

from __future__ import annotations

import re
import shutil
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.config.app_settings import EXCEL_TEMPLATES_DIR
from app.io_module.contracts import (
    CanonicalTransactionRow,
    CategorySummaryRow,
    CSVColumnMapping,
    ExportReport,
    ExternalSpreadsheetFormatterHook,
)
from app.io_module.csv_handler import (
    CSVSpreadsheetHandler,
    infer_transaction_kind,
    parse_flexible_amount,
    parse_flexible_date,
    parse_installment_info,
)
from app.io_module.duplicate_checker import normalize_text_for_hash

TEMPLATE_FILES: Dict[str, str] = {
    "despesas": "template_despesas_excel.xlsx",
    "receitas": "template_receitas_excel.xlsx",
    "transferencias": "template_transferencias_excel.xlsx",
    "extrato": "template_extrato_excel.xlsx",
    "categorias": "template_rel_group_cat_excel.xlsx",
    "investimento": "template_calculadora_investimento_excel.xlsx",
}

CURRENCY_FORMAT_BR = 'R$ #,##0.00;[Red]-R$ #,##0.00;"-"'
DATE_FORMAT_BR = "dd/mm/yyyy"


class ExcelSpreadsheetHandler:
    """Leitor e exportador modular de planilhas Excel (`.xlsx`)."""

    def __init__(
        self,
        templates_dir: Optional[Path] = None,
        external_hooks: Optional[List[ExternalSpreadsheetFormatterHook]] = None,
    ) -> None:
        self.templates_dir = Path(templates_dir) if templates_dir else EXCEL_TEMPLATES_DIR
        self._external_hooks: List[ExternalSpreadsheetFormatterHook] = list(
            external_hooks or []
        )

    def register_external_hook(self, hook: ExternalSpreadsheetFormatterHook) -> None:
        """Registra um formatador/pós-processador externo de planilhas Excel."""
        self._external_hooks.append(hook)

    # ------------------------------------------------------------------ #
    # LEITURA (.xlsx -> CanonicalTransactionRow)
    # ------------------------------------------------------------------ #
    def parse_file(
        self,
        file_path: Path | str,
        *,
        default_kind: str = "despesa",
        mapping: Optional[CSVColumnMapping] = None,
        sheet_name: Optional[str] = None,
    ) -> List[CanonicalTransactionRow]:
        """
        Importa linhas de uma planilha `.xlsx` (seja gerada pelos templates oficiais
        ou planilha bancária externa) convertendo-as em `CanonicalTransactionRow`.
        """
        path = Path(file_path)
        wb = openpyxl.load_workbook(path, data_only=True)
        sheets_to_read: List[Worksheet] = []
        if sheet_name and sheet_name in wb.sheetnames:
            sheets_to_read.append(wb[sheet_name])
        else:
            for sname in wb.sheetnames:
                norm_s = normalize_text_for_hash(sname)
                if norm_s in ("categorias", "simulacao_investimento", "resumo"):
                    continue
                sheets_to_read.append(wb[sname])
            if not sheets_to_read and wb.sheetnames:
                sheets_to_read.append(wb.active)

        all_rows: List[CanonicalTransactionRow] = []
        for ws in sheets_to_read:
            sname_norm = normalize_text_for_hash(ws.title)
            sheet_default_kind = default_kind
            if "receita" in sname_norm:
                sheet_default_kind = "receita"
            elif "transferencia" in sname_norm:
                sheet_default_kind = "transferencia"
            elif "despesa" in sname_norm:
                sheet_default_kind = "despesa"

            raw_matrix = list(ws.iter_rows(values_only=True))
            non_empty = [
                (idx + 1, list(r))
                for idx, r in enumerate(raw_matrix)
                if r and any(c is not None and str(c).strip() != "" for c in r)
            ]
            if not non_empty:
                continue

            first_row_strs = [str(c or "").strip() for c in non_empty[0][1]]
            header_map = CSVSpreadsheetHandler._infer_columns_from_header(first_row_strs)
            start_pos = 1 if header_map else 0

            c_data = (
                mapping.col_data
                if mapping and mapping.col_data is not None
                else header_map.get("data_vencimento", 0)
            )
            c_desc = (
                mapping.col_descricao
                if mapping and mapping.col_descricao is not None
                else header_map.get("descricao", 1)
            )
            c_val = (
                mapping.col_valor
                if mapping and mapping.col_valor is not None
                else header_map.get("valor", 2)
            )
            c_conta_cartao = header_map.get("conta_cartao")
            c_cat = header_map.get("categoria")
            c_sub = header_map.get("subcategoria")
            c_sit = header_map.get("situacao")
            c_conta = (
                mapping.col_conta
                if mapping and mapping.col_conta is not None
                else header_map.get("conta", c_conta_cartao)
            )
            c_dest = header_map.get("conta_destino")
            c_cartao = (
                mapping.col_cartao
                if mapping and mapping.col_cartao is not None
                else header_map.get("cartao", c_conta_cartao)
            )
            c_tags = header_map.get("tags")
            c_obs = header_map.get("observacoes")
            c_tipo = header_map.get("tipo")
            c_lanc = header_map.get("data_lancamento")
            c_efet = header_map.get("data_efetivacao")
            c_rec = header_map.get("recorrencia")

            for excel_row_num, row_vals in non_empty[start_pos:]:
                def cell_val(col_i: Optional[int]) -> Any:
                    if col_i is None or col_i < 0 or col_i >= len(row_vals):
                        return None
                    return row_vals[col_i]

                def cell_str(col_i: Optional[int]) -> str:
                    v = cell_val(col_i)
                    return "" if v is None else str(v).strip()

                dt_venc = parse_flexible_date(cell_val(c_data))
                if dt_venc is None:
                    for cand in (c_lanc, c_efet, 0, 1, 5, 6):
                        dt_venc = parse_flexible_date(cell_val(cand))
                        if dt_venc is not None:
                            break
                if dt_venc is None:
                    continue

                val_abs, is_neg = parse_flexible_amount(cell_val(c_val))
                if val_abs <= 0:
                    continue

                desc = cell_str(c_desc) or "Lançamento Importado"
                has_dest = bool(c_dest is not None and cell_str(c_dest))
                kind = infer_transaction_kind(
                    tipo_raw=cell_str(c_tipo),
                    descricao=desc,
                    categoria=cell_str(c_cat),
                    is_neg=is_neg,
                    neg_is_expense=True if mapping is None else bool(mapping.valor_negativo_despesas),
                    default_kind=sheet_default_kind,
                    has_destination=has_dest,
                )

                sit_norm = normalize_text_for_hash(cell_str(c_sit))
                efetivada = (
                    sit_norm
                    in (
                        "pago",
                        "paga",
                        "recebido",
                        "recebida",
                        "efetivado",
                        "efetivada",
                        "pg",
                        "s",
                        "sim",
                    )
                    if c_sit is not None
                    else True
                )
                tags_raw = cell_str(c_tags)
                tags = (
                    [t.strip() for t in re.split(r"[,;|]", tags_raw) if t.strip()]
                    if tags_raw
                    else []
                )

                num_parc, total_parcs, rec_val = parse_installment_info(
                    cell_val(c_rec), desc
                )

                all_rows.append(
                    CanonicalTransactionRow(
                        kind=kind,
                        descricao=desc,
                        valor=val_abs,
                        data_vencimento=dt_venc,
                        data_lancamento=parse_flexible_date(cell_val(c_lanc)) or dt_venc,
                        data_efetivacao=parse_flexible_date(cell_val(c_efet))
                        or (dt_venc if efetivada else None),
                        efetivada=efetivada,
                        categoria=cell_str(c_cat) or "Outros",
                        subcategoria=cell_str(c_sub) or "Outros",
                        conta=cell_str(c_conta),
                        conta_destino=cell_str(c_dest),
                        cartao=cell_str(c_cartao),
                        recorrencia=rec_val,
                        tags=tags,
                        observacoes=cell_str(c_obs),
                        raw_row_index=excel_row_num,
                        num_parcela=num_parc,
                        num_parcelas=total_parcs,
                        is_negative=is_neg,
                    )
                )
        wb.close()
        return all_rows

    # ------------------------------------------------------------------ #
    # EXPORTAÇÃO (CanonicalTransactionRow -> .xlsx)
    # ------------------------------------------------------------------ #
    @staticmethod
    def _style_worksheet(ws: Worksheet, currency_cols: Sequence[int], date_cols: Sequence[int]) -> None:
        header_fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
        header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
        body_font = Font(name="Segoe UI", size=10)
        thin_border = Border(
            bottom=Side(style="thin", color="E2E8F0"),
        )

        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="left", vertical="center")

        for r_idx in range(2, ws.max_row + 1):
            for c_idx in range(1, ws.max_column + 1):
                c = ws.cell(row=r_idx, column=c_idx)
                c.font = body_font
                c.border = thin_border
                if c_idx in currency_cols:
                    c.number_format = CURRENCY_FORMAT_BR
                    c.alignment = Alignment(horizontal="right", vertical="center")
                elif c_idx in date_cols:
                    c.number_format = DATE_FORMAT_BR
                    c.alignment = Alignment(horizontal="center", vertical="center")

        for col_idx in range(1, ws.max_column + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 10
            for row_idx in range(1, min(ws.max_row + 1, 200)):
                val = ws.cell(row=row_idx, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, min(len(str(val)) + 4, 48))
            ws.column_dimensions[col_letter].width = max_len

    def _load_or_create_template(self, mode_norm: str, target_path: Path) -> Tuple[Workbook, Optional[str]]:
        tpl_name = TEMPLATE_FILES.get(mode_norm)
        if tpl_name:
            tpl_path = self.templates_dir / tpl_name
            if tpl_path.exists():
                shutil.copy2(tpl_path, target_path)
                return openpyxl.load_workbook(target_path), tpl_name
        wb = openpyxl.Workbook()
        return wb, None

    @staticmethod
    def _clear_template_placeholder_rows(ws: Worksheet) -> None:
        """Remove linhas em branco pré-existentes abaixo do cabeçalho (linha 1) nos templates."""
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)

    def _populate_despesas_sheet(self, ws: Worksheet, rows: Sequence[CanonicalTransactionRow]) -> None:
        if ws.max_row < 1 or ws.cell(1, 1).value is None:
            ws.append(
                [
                    "Descrição",
                    "Valor",
                    "Situação",
                    "Categoria",
                    "Subcategoria",
                    "Lançamento",
                    "Vencimento",
                    "Efetivação",
                    "Recorrência",
                    "Conta",
                    "Cartão",
                    "Tags",
                    "Observações",
                ]
            )
        else:
            self._clear_template_placeholder_rows(ws)

        for r in rows:
            ws.append(
                [
                    r.descricao,
                    round(float(r.valor), 2),
                    "Pago" if r.efetivada else "Pendente",
                    r.categoria,
                    r.subcategoria,
                    (r.data_lancamento or r.data_vencimento).strftime("%d/%m/%Y"),
                    r.data_vencimento.strftime("%d/%m/%Y"),
                    r.data_efetivacao.strftime("%d/%m/%Y") if r.data_efetivacao else "",
                    r.recorrencia,
                    r.conta,
                    r.cartao,
                    ", ".join(r.tags),
                    r.observacoes,
                ]
            )
        self._style_worksheet(ws, currency_cols=[2], date_cols=[6, 7, 8])

    def _populate_receitas_sheet(self, ws: Worksheet, rows: Sequence[CanonicalTransactionRow]) -> None:
        if ws.max_row < 1 or ws.cell(1, 1).value is None:
            ws.append(
                [
                    "Descrição",
                    "Valor",
                    "Situação",
                    "Categoria",
                    "Subcategoria",
                    "Lançamento",
                    "Vencimento",
                    "Efetivação",
                    "Recorrência",
                    "Conta",
                    "Tags",
                    "Observações",
                ]
            )
        else:
            self._clear_template_placeholder_rows(ws)

        for r in rows:
            ws.append(
                [
                    r.descricao,
                    round(float(r.valor), 2),
                    "Recebido" if r.efetivada else "Pendente",
                    r.categoria,
                    r.subcategoria,
                    (r.data_lancamento or r.data_vencimento).strftime("%d/%m/%Y"),
                    r.data_vencimento.strftime("%d/%m/%Y"),
                    r.data_efetivacao.strftime("%d/%m/%Y") if r.data_efetivacao else "",
                    r.recorrencia,
                    r.conta,
                    ", ".join(r.tags),
                    r.observacoes,
                ]
            )
        self._style_worksheet(ws, currency_cols=[2], date_cols=[6, 7, 8])

    def _populate_transferencias_sheet(self, ws: Worksheet, rows: Sequence[CanonicalTransactionRow]) -> None:
        if ws.max_row < 1 or ws.cell(1, 1).value is None:
            ws.append(
                [
                    "Descrição",
                    "Vencimento",
                    "Origem",
                    "Destino",
                    "Valor",
                    "Tags",
                    "Observações",
                ]
            )
        else:
            self._clear_template_placeholder_rows(ws)

        for r in rows:
            ws.append(
                [
                    r.descricao,
                    r.data_vencimento.strftime("%d/%m/%Y"),
                    r.conta,
                    r.conta_destino,
                    round(float(r.valor), 2),
                    ", ".join(r.tags),
                    r.observacoes,
                ]
            )
        self._style_worksheet(ws, currency_cols=[5], date_cols=[2])

    def _populate_extrato_sheet(self, ws: Worksheet, rows: Sequence[CanonicalTransactionRow]) -> None:
        if ws.max_row < 1 or ws.cell(1, 1).value is None:
            ws.append(["Data", "Descrição", "Categoria", "Subcategoria", "Conta", "Valor"])
        else:
            self._clear_template_placeholder_rows(ws)

        for r in rows:
            signed_val = -round(abs(r.valor), 2) if r.kind == "despesa" else round(abs(r.valor), 2)
            ws.append(
                [
                    r.data_vencimento.strftime("%d/%m/%Y"),
                    r.descricao,
                    r.categoria,
                    r.subcategoria,
                    r.conta or r.cartao,
                    signed_val,
                ]
            )
        self._style_worksheet(ws, currency_cols=[6], date_cols=[1])

    def _populate_categorias_sheet(self, ws: Worksheet, summary: Sequence[CategorySummaryRow]) -> None:
        if ws.max_row < 1 or ws.cell(1, 1).value is None:
            ws.append(["Categoria", "Total"])
        else:
            self._clear_template_placeholder_rows(ws)

        for cs in summary:
            ws.append([cs.categoria, round(float(cs.total), 2)])
        self._style_worksheet(ws, currency_cols=[2], date_cols=[])

    def export_transactions(
        self,
        file_path: Path | str,
        rows: Sequence[CanonicalTransactionRow],
        *,
        mode: str = "extrato",
        category_summary: Optional[Sequence[CategorySummaryRow]] = None,
        external_hooks: Optional[Sequence[ExternalSpreadsheetFormatterHook]] = None,
    ) -> ExportReport:
        """
        Exporta lançamentos para `.xlsx` utilizando os templates oficiais de `excel_templates/`
        ou gerando uma pasta de trabalho multi-abas (`mode='completo'`).
        """
        path = Path(file_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        mode_norm = (mode or "extrato").lower()

        if mode_norm == "completo":
            wb = openpyxl.Workbook()
            ws_ext = wb.active
            ws_ext.title = "Extrato"
            self._populate_extrato_sheet(ws_ext, rows)

            desp = [r for r in rows if r.kind == "despesa"]
            rec = [r for r in rows if r.kind == "receita"]
            transf = [r for r in rows if r.kind == "transferencia"]

            ws_d = wb.create_sheet("Despesas")
            self._populate_despesas_sheet(ws_d, desp)

            ws_r = wb.create_sheet("Receitas")
            self._populate_receitas_sheet(ws_r, rec)

            ws_t = wb.create_sheet("Transferencias")
            self._populate_transferencias_sheet(ws_t, transf)

            if category_summary is not None:
                ws_c = wb.create_sheet("Categorias")
                self._populate_categorias_sheet(ws_c, category_summary)

            wb.save(path)
            wb.close()
            exported_count = len(rows)
            tpl_used = "multi_sheet_workbook"
        else:
            wb, tpl_used = self._load_or_create_template(mode_norm, path)
            ws = wb.active
            if mode_norm == "despesas":
                self._populate_despesas_sheet(ws, rows)
                exported_count = len(rows)
            elif mode_norm == "receitas":
                self._populate_receitas_sheet(ws, rows)
                exported_count = len(rows)
            elif mode_norm == "transferencias":
                self._populate_transferencias_sheet(ws, rows)
                exported_count = len(rows)
            elif mode_norm == "categorias" and category_summary is not None:
                self._populate_categorias_sheet(ws, category_summary)
                exported_count = len(category_summary)
            else:
                self._populate_extrato_sheet(ws, rows)
                exported_count = len(rows)
            wb.save(path)
            wb.close()

        applied_hooks: List[str] = []
        all_hooks = [*self._external_hooks, *(external_hooks or [])]
        for hook in all_hooks:
            hook_name = getattr(hook, "__name__", hook.__class__.__name__)
            hook(
                path,
                {
                    "format": "xlsx",
                    "mode": mode_norm,
                    "rows": exported_count,
                    "template": tpl_used,
                },
            )
            applied_hooks.append(hook_name)

        return ExportReport(
            target_path=str(path),
            format_type="xlsx",
            exported_rows=exported_count,
            template_used=tpl_used,
            external_formatters_applied=applied_hooks,
        )
