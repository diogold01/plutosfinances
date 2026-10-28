"""
Testes Automatizados da Etapa 6/6:
- Módulo Independente de Importação e Exportação CSV e Excel (`app/io_module/`)
- Hooks para ferramentas externas de formatação de planilhas (`ExternalSpreadsheetFormatterHook`)
- Prevenção de duplicidade via SHA-256 (`DuplicateChecker` & `import_fingerprints`)
- Persistência de modelos CSV (`modelos_csv` e `modelos_csv_campos`)
- Integração End-to-End com `FerramentasView`, `MainWindow` e `BackupService`
"""

from __future__ import annotations

import os
from datetime import date
from pathlib import Path
from typing import Any, Dict

import openpyxl
from PySide6.QtWidgets import QApplication

from app.config.app_settings import AppSettings
from app.database.backup_service import BackupSchedulerService
from app.database.connection_manager import ConnectionManager
from app.database.repositories.account_repository import AccountRepository
from app.domain.models import KIND_DESPESA, KIND_RECEITA
from app.domain.services.transaction_service import TransactionService
from app.io_module import (
    CanonicalTransactionRow,
    CategorySummaryRow,
    CSVColumnMapping,
    CSVSpreadsheetHandler,
    DataIOOrchestrator,
    DuplicateChecker,
    ExcelSpreadsheetHandler,
)
from app.ui.main_window import MainWindow

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def _get_qapp() -> QApplication:
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


def test_decoupled_csv_and_excel_handlers_with_external_formatter_hook(tmp_path: Path) -> None:
    """Valida a independência do módulo de I/O (sem banco/UI) e o hook para formatadores externos."""
    sample_rows = [
        CanonicalTransactionRow(
            kind="despesa",
            descricao="Supermercado Pão de Açúcar",
            valor=349.90,
            data_vencimento=date(2026, 10, 10),
            data_lancamento=date(2026, 10, 10),
            data_efetivacao=date(2026, 10, 10),
            efetivada=True,
            categoria="Alimentação",
            subcategoria="Supermercado",
            conta="Nubank",
            tags=["essencial", "casa"],
            observacoes="Compra do mês",
        ),
        CanonicalTransactionRow(
            kind="receita",
            descricao="Consultoria Técnica",
            valor=4200.00,
            data_vencimento=date(2026, 10, 15),
            data_lancamento=date(2026, 10, 15),
            data_efetivacao=date(2026, 10, 15),
            efetivada=True,
            categoria="Salário",
            subcategoria="Outros",
            conta="Itaú",
            tags=["trabalho"],
            observacoes="NF 102",
        ),
    ]

    hook_invocations: list[Dict[str, Any]] = []

    def external_corporate_styler(file_path: Path, metadata: Dict[str, Any]) -> None:
        hook_invocations.append({"path": str(file_path), **metadata})
        if metadata["format"] == "xlsx":
            wb = openpyxl.load_workbook(file_path)
            wb.active.sheet_properties.tabColor = "1072BA"
            wb.save(file_path)
            wb.close()

    # 1. Exportação e reimportação CSV independente
    csv_handler = CSVSpreadsheetHandler(external_hooks=[external_corporate_styler])
    csv_path = tmp_path / "extrato_teste.csv"
    exp_csv = csv_handler.export_transactions(csv_path, sample_rows, mode="extrato")
    assert exp_csv.exported_rows == 2
    assert "external_corporate_styler" in exp_csv.external_formatters_applied

    parsed_csv = csv_handler.parse_file(csv_path)
    assert len(parsed_csv) == 2
    assert parsed_csv[0].descricao == "Supermercado Pão de Açúcar"
    assert parsed_csv[0].valor == 349.90
    assert parsed_csv[0].kind == "despesa"
    assert parsed_csv[1].descricao == "Consultoria Técnica"
    assert parsed_csv[1].valor == 4200.00
    assert parsed_csv[1].kind == "receita"

    # 2. Exportação e reimportação Excel (.xlsx) usando templates oficiais e hook externo
    excel_handler = ExcelSpreadsheetHandler(external_hooks=[external_corporate_styler])
    xlsx_desp_path = tmp_path / "despesas_template.xlsx"
    exp_xlsx = excel_handler.export_transactions(
        xlsx_desp_path, [sample_rows[0]], mode="despesas"
    )
    assert exp_xlsx.exported_rows == 1
    assert exp_xlsx.template_used == "template_despesas_excel.xlsx"
    assert "external_corporate_styler" in exp_xlsx.external_formatters_applied

    parsed_xlsx = excel_handler.parse_file(xlsx_desp_path, default_kind="despesa")
    assert len(parsed_xlsx) == 1
    assert parsed_xlsx[0].descricao == "Supermercado Pão de Açúcar"
    assert parsed_xlsx[0].valor == 349.90
    assert parsed_xlsx[0].categoria == "Alimentação"
    assert len(hook_invocations) == 2


def test_orchestrator_import_export_deduplication_and_csv_mappings(tmp_path: Path) -> None:
    """Valida importação/exportação com banco SQLite, prevenção de duplicidade SHA-256 e `modelos_csv`."""
    db_path = tmp_path / "io_test.sqlite"
    settings = AppSettings.load(tmp_path / "settings.json")
    settings.active_db_path = str(db_path)
    settings.backup_dir = str(tmp_path / "backups")
    settings.save()
    db = ConnectionManager(settings=settings, auto_connect=False)
    db.create_new_database(db_path, overwrite=True)

    acc_repo = AccountRepository(db)
    acc_id = acc_repo.create_account("Conta Inter", saldo_inicial=1000.0)

    orchestrator = DataIOOrchestrator(db)

    # 1. Persistência de modelo CSV em `modelos_csv` e `modelos_csv_campos`
    mapping = CSVColumnMapping(
        descricao_modelo="Modelo Banco Inter CSV",
        separador=";",
        possui_cabecalho=True,
        formato_data="dd/MM/yyyy",
        valor_negativo_despesas=True,
        id_conta=acc_id,
        col_data=0,
        col_descricao=1,
        col_valor=2,
        col_categoria=3,
    )
    saved_id = orchestrator.save_csv_mapping(mapping)
    loaded_mappings = orchestrator.list_csv_mappings()
    assert any(m.id == saved_id and m.col_valor == 2 for m in loaded_mappings)

    # 2. Cria um CSV de extrato bancário e importa pela 1ª vez
    raw_csv = tmp_path / "extrato_inter.csv"
    raw_csv.write_text(
        "Data;Descrição;Valor;Categoria\n"
        "12/10/2026;Padaria Central;-45,50;Alimentação\n"
        "14/10/2026;Reembolso Cliente;320,00;Salário\n"
        "16/10/2026; Posto Shell ;-210,00;Transporte\n",
        encoding="utf-8-sig",
    )

    rep1 = orchestrator.import_file(
        raw_csv,
        profile_id="1",
        target_account_id=acc_id,
        mapping=mapping,
        skip_duplicates=True,
    )
    assert rep1.total_rows_read == 3
    assert rep1.imported_count == 3
    assert rep1.skipped_duplicates == 0
    assert rep1.error_count == 0

    # 3. Reimporta o mesmo arquivo CSV -> todos os 3 lançamentos devem ser barrados pelo DuplicateChecker!
    rep2 = orchestrator.import_file(
        raw_csv,
        profile_id="1",
        target_account_id=acc_id,
        mapping=mapping,
        skip_duplicates=True,
    )
    assert rep2.total_rows_read == 3
    assert rep2.imported_count == 0
    assert rep2.skipped_duplicates == 3

    # 4. Exporta pasta de trabalho completa multi-abas (.xlsx) e resumo por categorias (.xlsx)
    full_xlsx = tmp_path / "export_completo_2026_10.xlsx"
    exp_full = orchestrator.export_month_data(
        full_xlsx, year=2026, month=10, export_mode="completo"
    )
    assert exp_full.exported_rows == 3
    wb = openpyxl.load_workbook(full_xlsx)
    assert {"Extrato", "Despesas", "Receitas", "Transferencias", "Categorias"}.issubset(
        set(wb.sheetnames)
    )
    wb.close()

    db.close()


def test_end_to_end_ui_ferramentas_io_and_backup_rotation(tmp_path: Path) -> None:
    """Teste End-to-End integrando MainWindow, FerramentasView (Import/Export) e BackupService."""
    _get_qapp()
    db_path = tmp_path / "e2e_financas.sqlite"
    backup_dir = tmp_path / "backups_e2e"
    settings = AppSettings.load(tmp_path / "settings_e2e.json")
    settings.active_db_path = str(db_path)
    settings.backup_dir = str(backup_dir)
    settings.save()
    db = ConnectionManager(settings=settings, auto_connect=False)
    db.create_new_database(db_path, overwrite=True)
    backup = BackupSchedulerService(db, settings)

    win = MainWindow(db, backup)
    win.current_year = 2026
    win.current_month = 10

    # Navega até a tela Ferramentas e testa importação/exportação via UI
    win.navigate_to("ferramentas")
    ferramentas = win.ferramentas_view

    csv_in = tmp_path / "import_ui.csv"
    csv_in.write_text(
        "Data;Descrição;Tipo;Categoria;Subcategoria;Valor\n"
        "05/10/2026;Assinatura Fibra;Despesa;Moradia;Internet;-129,90\n"
        "08/10/2026;Dividendos FII;Receita;Investimentos;FII;450,00\n",
        encoding="utf-8-sig",
    )
    imp_rep = ferramentas.execute_import_from_path(csv_in)
    assert imp_rep.imported_count == 2

    xlsx_out = tmp_path / "export_ui.xlsx"
    exp_rep = ferramentas.execute_export_to_path(xlsx_out)
    assert exp_rep.exported_rows >= 2
    assert xlsx_out.exists()

    # Verifica que a mutação gerada pela importação aciona o backup incremental online do SQLite
    b_res = backup.run_backup_cycle()
    assert b_res.performed and b_res.created_file is not None and b_res.created_file.exists()

    win.close()
    db.close()


def test_ai_prompt_generator_and_notification_bell_in_main_window(tmp_path: Path) -> None:
    """Valida o gerador de prompt para IA em FerramentasView e a Central de Notificações com badge em MainWindow."""
    from datetime import date
    _get_qapp()

    db_path = tmp_path / "notif_test.sqlite"
    backup_dir = tmp_path / "backups_notif"
    settings = AppSettings.load(tmp_path / "settings_notif.json")
    settings.active_db_path = str(db_path)
    settings.backup_dir = str(backup_dir)
    settings.save()

    db = ConnectionManager(settings=settings, auto_connect=False)
    db.create_new_database(db_path, overwrite=True)
    backup = BackupSchedulerService(db, settings)

    win = MainWindow(db, backup)
    try:
        # 1. Valida o Gerador de Prompt para IA (FerramentasView)
        ferramentas = win.ferramentas_view
        prompt = ferramentas.build_ai_import_prompt()
        assert "data;descricao;valor;categoria;subcategoria;conta_cartao;parcela;observacao" in prompt
        assert "Plutos" in prompt
        assert "dd/MM/yyyy" in prompt
        assert "[Terceiro: NomeDaPessoa]" in prompt

        # 2. Valida o Sininho de Notificações no Header
        assert hasattr(win, "btn_notifications")
        assert not win.btn_notifications.icon().isNull() or "🔔" in win.btn_notifications.text()

        # Cria uma despesa a pagar com vencimento hoje
        today = date.today()
        win.tx_service.create_expense(
            descricao="Energia Elétrica",
            valor=185.50,
            data_vencimento=today,
            pago=False,
        )

        # Atualiza as notificações e verifica o badge counter
        win._check_system_notifications()
        assert len(win._unread_notifications) >= 1
        assert any("Conta a Pagar (hoje)" in n[1] for n in win._unread_notifications)
        assert str(len(win._unread_notifications)) in win.btn_notifications.text()

        # Valida que o método de abertura de menu roda sem falhas
        win._show_notifications_menu()
    finally:
        win.close()
        db.close()


def test_import_with_installments_sets_recorrencia_p_and_labels(tmp_path: Path) -> None:
    """Valida que importação de CSV com coluna parcela ou texto extrai parcelas e salva recorrencia='P'."""
    db_path = tmp_path / "test_parc.sqlite"
    settings = AppSettings.load(tmp_path / "settings_parc.json")
    settings.active_db_path = str(db_path)
    settings.save()

    db = ConnectionManager(settings=settings, auto_connect=False)
    db.create_new_database(db_path, overwrite=True)

    csv_content = (
        "data;descricao;valor;categoria;subcategoria;conta_cartao;parcela;observacao\n"
        "02/10/2026;Alura;97,66;Educação;Outros;;8/12;\n"
        "02/10/2026;MercadoLivre - Parcela 2/6;41,65;Outros;Outros;;;\n"
        "15/10/2026;Google One;35,99;Serviços;Outros;;1/1;\n"
    )
    csv_file = tmp_path / "fatura_parcelas.csv"
    csv_file.write_text(csv_content, encoding="utf-8")

    orchestrator = DataIOOrchestrator(db)
    rep = orchestrator.import_file(csv_file, profile_id="1")
    assert rep.imported_count == 3
    assert rep.error_count == 0

    # Consulta despesas no banco
    rows = db.execute_query(
        "SELECT descricao, recorrencia, num_parcela, num_parcelas FROM despesas ORDER BY id ASC"
    )
    by_desc = {r["descricao"]: r for r in rows}

    assert by_desc["Alura"]["recorrencia"] == "P"
    assert by_desc["Alura"]["num_parcela"] == 8
    assert by_desc["Alura"]["num_parcelas"] == 12

    assert by_desc["MercadoLivre - Parcela 2/6"]["recorrencia"] == "P"
    assert by_desc["MercadoLivre - Parcela 2/6"]["num_parcela"] == 2
    assert by_desc["MercadoLivre - Parcela 2/6"]["num_parcelas"] == 6

    assert by_desc["Google One"]["recorrencia"] == "U"
    assert by_desc["Google One"]["num_parcela"] == 1
    assert by_desc["Google One"]["num_parcelas"] == 1

    # Valida labels de ocorrências via TransactionService
    tx = TransactionService(db)
    occs = tx.list_occurrences(2026, 10, profile_id="1")
    occ_map = {o.descricao: o for o in occs}
    assert occ_map["Alura"].parcela_label == "8/12"
    assert occ_map["MercadoLivre - Parcela 2/6"].parcela_label == "2/6"
    assert (occ_map["Google One"].parcela_label or "À vista") == "À vista"

    db.close()


def test_automatic_income_expense_detection_from_positive_values_and_excel_dates(tmp_path: Path) -> None:
    """
    Valida o caso real do usuário:
    - Planilha gerada por IA com 8 colunas (sem coluna 'tipo' explícita)
    - Todos os valores positivos (inclusive Salário e Transferência Recebida)
    - Coluna parcela corrompida pelo Excel como '01/jan'
    - Coluna 'conta_cartao' unificada com 'Nubank'
    Garante que Salário e Transferências Recebidas entram como RECEITA e compras como DESPESA.
    """
    db_path = tmp_path / "test_auto_kind.sqlite"
    settings = AppSettings.load(tmp_path / "settings_auto.json")
    settings.active_db_path = str(db_path)
    settings.save()

    db = ConnectionManager(settings=settings, auto_connect=False)
    db.create_new_database(db_path, overwrite=True)

    csv_data = (
        "data;descricao;valor;categoria;subcategoria;conta_cartao;parcela;observacao\n"
        "03/10/2026;Transferência Recebida - Nicolas Mantovani Marq;10;Outros;Outros;Nubank;01/jan;\n"
        "06/10/2026;Salário - Coliseu Tecnologia e Consultoria Ltda;1900;Outros;Outros;Nubank;01/jan;\n"
        "07/10/2026;Pagamento de Fatura;515,99;Pagamentos;Outros;Nubank;01/jan;\n"
        "07/10/2026;Auto Posto Xará;15,09;Transporte;Combustível;Nubank;01/jan;\n"
        "07/10/2026;Império Lanches;60;Alimentação;Lanche;Nubank;01/jan;\n"
        "07/10/2026;Transferência Nathaly Beatriz Leite;540;Moradia;Fixas;Nubank;01/jan;[Terceiro: Nathaly]\n"
    )
    csv_file = tmp_path / "extrato_plutos_outubro_2026.csv"
    csv_file.write_text(csv_data, encoding="utf-8-sig")

    orchestrator = DataIOOrchestrator(db)
    rep = orchestrator.import_file(csv_file, profile_id="1")
    assert rep.imported_count == 6
    assert rep.error_count == 0

    # 1. Verifica no banco se as entradas viraram receitas e as saídas viraram despesas
    receitas = db.execute_query("SELECT descricao, valor FROM receitas ORDER BY id ASC")
    despesas = db.execute_query("SELECT descricao, valor FROM despesas ORDER BY id ASC")

    rec_descs = [r["descricao"] for r in receitas]
    desp_descs = [d["descricao"] for d in despesas]

    assert "Salário - Coliseu Tecnologia e Consultoria Ltda" in rec_descs
    assert "Transferência Recebida - Nicolas Mantovani Marq" in rec_descs
    assert len(receitas) == 2

    assert "Pagamento de Fatura" in desp_descs
    assert "Auto Posto Xará" in desp_descs
    assert "Império Lanches" in desp_descs
    assert "Transferência Nathaly Beatriz Leite" in desp_descs
    assert len(despesas) == 4

    # 2. Verifica as ocorrências no TransactionService
    tx = TransactionService(db)
    occs = tx.list_occurrences(2026, 10, profile_id="1")
    by_desc = {o.descricao: o for o in occs}

    assert by_desc["Salário - Coliseu Tecnologia e Consultoria Ltda"].kind == "RECEITA"
    assert by_desc["Salário - Coliseu Tecnologia e Consultoria Ltda"].valor == 1900.0
    assert by_desc["Salário - Coliseu Tecnologia e Consultoria Ltda"].recorrencia == "U"
    assert (by_desc["Salário - Coliseu Tecnologia e Consultoria Ltda"].parcela_label or "Única") in ("Única", "À vista")

    assert by_desc["Transferência Recebida - Nicolas Mantovani Marq"].kind == "RECEITA"
    assert by_desc["Transferência Recebida - Nicolas Mantovani Marq"].valor == 10.0

    assert by_desc["Pagamento de Fatura"].kind == "DESPESA"
    assert by_desc["Pagamento de Fatura"].valor == 515.99

    db.close()


def test_e2e_import_and_conversion_affects_balance_and_summaries(tmp_path: Path) -> None:
    """
    Etapa 3 - Validação Completa Ponta a Ponta:
    1. Importa extrato real com classificação automática de receitas e despesas.
    2. Efetiva lançamentos e calcula saldos e indicadores do mês.
    3. Converte Despesa em Receita via TransactionService.update_occurrence (novo tipo)
       e valida o recálculo imediato do saldo da conta bancária (+2x valor) e dos totais.
    4. Converte Receita em Despesa e valida o estorno proporcional nos indicadores.
    """
    db_path = tmp_path / "test_e2e_balance.sqlite"
    settings = AppSettings.load(tmp_path / "settings_e2e.json")
    settings.active_db_path = str(db_path)
    settings.save()

    db = ConnectionManager(settings=settings, auto_connect=False)
    db.create_new_database(db_path, overwrite=True)

    acc_repo = AccountRepository(db)
    # Garante que a conta 'Nubank' existe com saldo inicial de R$ 1000,00
    acc_id = acc_repo.create_account(
        descricao="Nubank",
        saldo_inicial=1000.0,
        mes_inicial=10,
        ano_inicial=2026,
    )

    csv_data = (
        "data;descricao;valor;categoria;subcategoria;conta_cartao;parcela;observacao\n"
        "03/10/2026;Transferência Recebida - Nicolas Mantovani Marq;10;Outros;Outros;Nubank;01/jan;\n"
        "06/10/2026;Salário - Coliseu Tecnologia e Consultoria Ltda;1900;Outros;Outros;Nubank;01/jan;\n"
        "07/10/2026;Auto Posto Xará;50;Transporte;Combustível;Nubank;01/jan;\n"
    )
    csv_file = tmp_path / "extrato_e2e.csv"
    csv_file.write_text(csv_data, encoding="utf-8-sig")

    orchestrator = DataIOOrchestrator(db)
    rep = orchestrator.import_file(csv_file, profile_id="1")
    assert rep.imported_count == 3
    assert rep.error_count == 0

    tx = TransactionService(db)
    occs = tx.list_occurrences(2026, 10, profile_id="1")
    by_desc = {o.descricao: o for o in occs}

    # Efetiva todas as ocorrências
    for o in occs:
        tx.set_effective(o, True)

    # Recarrega ocorrências efetivadas
    occs_efet = tx.list_occurrences(2026, 10, profile_id="1")
    resumo_inicial = tx.summarize(occs_efet)
    assert resumo_inicial["receitas_efetivadas"] == 1910.00  # 1900 + 10
    assert resumo_inicial["despesas_pagas"] == 50.00     # 50

    # Saldo da conta: Inicial (1000) + Receitas (1910) - Despesas (50) = 2860.00
    bal_inicial = acc_repo.get_account_summary(acc_id, 10, 2026)
    assert bal_inicial.saldo_atual == 2860.00

    # 3. CONVERTE A DESPESA 'Auto Posto Xará' (R$ 50) PARA RECEITA (Reembolso)
    desp_occ = [o for o in occs_efet if "Auto Posto Xará" in o.descricao][0]
    tx.update_occurrence(
        desp_occ,
        new_kind=KIND_RECEITA,
        descricao="Reembolso Combustível - Auto Posto Xará",
        valor=50.00,
        efetivada=True,
    )

    # Recalcula indicadores do mês
    occs_pos_conv = tx.list_occurrences(2026, 10, profile_id="1")
    resumo_pos_conv = tx.summarize(occs_pos_conv)
    assert resumo_pos_conv["receitas_efetivadas"] == 1960.00  # 1910 + 50
    assert resumo_pos_conv["despesas_pagas"] == 0.00      # 50 virou 0

    # Saldo da conta Nubank deve ter aumentado exatamente em 2 * 50 = 100:
    # 2860.00 + 100.00 = 2960.00 (1000 + 1960 - 0)
    bal_pos_conv = acc_repo.get_account_summary(acc_id, 10, 2026)
    assert bal_pos_conv.saldo_atual == 2960.00

    # 4. CONVERTE A RECEITA 'Transferência Recebida' (R$ 10) PARA DESPESA
    rec_occ = [o for o in occs_pos_conv if "Transferência Recebida" in o.descricao][0]
    tx.update_occurrence(
        rec_occ,
        new_kind=KIND_DESPESA,
        descricao="Transferência Enviada - Nicolas Mantovani Marq",
        valor=10.00,
        efetivada=True,
    )

    occs_final = tx.list_occurrences(2026, 10, profile_id="1")
    resumo_final = tx.summarize(occs_final)
    assert resumo_final["receitas_efetivadas"] == 1950.00  # 1960 - 10
    assert resumo_final["despesas_pagas"] == 10.00    # 0 + 10

    # Saldo da conta final: 1000 + 1950 - 10 = 2940.00
    bal_final = acc_repo.get_account_summary(acc_id, 10, 2026)
    assert bal_final.saldo_atual == 2940.00

    db.close()


def test_nubank_credit_card_invoice_import_closing_day_and_payment_filter(tmp_path: Path) -> None:
    """
    Valida a Etapa 1 (fatura real Nubank_2026-11-09.csv):
    1. Cartão com fechamento dia 02 e vencimento dia 09.
    2. Compras/parcelas faturadas no dia do fechamento (02/10/2026) vão para a fatura 11/2026 (09/11/2026)
       e NÃO vazam para a fatura 10/2026.
    3. A linha 'Pagamento recebido' ('- 515,99') é ignorada na importação da fatura.
    4. Estornos legítimos no cartão abatem (valor negativo) o total da fatura.
    """
    from app.domain.services.credit_card_service import CreditCardService

    db_path = tmp_path / "test_nubank_inv.sqlite"
    settings = AppSettings.load(tmp_path / "settings_nubank_inv.json")
    settings.active_db_path = str(db_path)
    settings.save()

    db = ConnectionManager(settings=settings, auto_connect=False)
    db.create_new_database(db_path, overwrite=True)

    acc_repo = AccountRepository(db)
    tx = TransactionService(db)
    cc = CreditCardService(db, tx)

    acc_id = acc_repo.create_account(descricao="Nubank Conta", saldo_inicial=2000.0)
    card = cc.create_card(
        nome="Nubank",
        limite=1750.0,
        dia_fecha=2,
        dia_venc=9,
        id_conta=acc_id,
        brand_key="master",
    )

    csv_content = (
        "date,title,amount\n"
        '2026-10-20,Mercado*Mercadolivre - Parcela 1/4,"77,25"\n'
        '2026-10-15,Google One Gemini,"35,99"\n'
        '2026-10-07,Pagamento recebido,"- 515,99"\n'
        '2026-10-06,Plano NuCel,"45,00"\n'
        '2026-10-05,IOF de compra internacional,"0,44"\n'
        '2026-10-05,Google Youtube,"16,90"\n'
        '2026-10-05,Wl *Steam Purchase,"12,82"\n'
        '2026-10-04,Pag*Steam,"23,99"\n'
        '2026-10-02,Dl *Aliexpress Br Alip - Parcela 3/3,"80,39"\n'
        '2026-10-02,Mercadolivre*Maxcostor - Parcela 2/6,"41,65"\n'
        '2026-10-02,Pag*Steam - Parcela 5/6,"26,76"\n'
        '2026-10-02,Alura - Parcela 8/12,"97,66"\n'
        '2026-10-02,Pag*Steam - Parcela 2/3,"6,16"\n'
        '2026-10-02,Boleto no Crédito - SOCIEDADE EDUCACIONAL LEONARDO DA VINCIS S LTDA - 3/3,"88,68"\n'
        '2026-10-02,Mp *Pichau - Parcela 5/8,"15,72"\n'
        '2026-10-02,Pag*Steam - Parcela 3/3,"13,33"\n'
    )
    csv_file = tmp_path / "Nubank_2026-11-09 (2).csv"
    csv_file.write_text(csv_content, encoding="utf-8")

    orchestrator = DataIOOrchestrator(db, tx)
    rep = orchestrator.import_file(
        csv_file,
        profile_id="1",
        target_card_id=card.id,
    )

    assert rep.total_rows_read == 16
    assert rep.skipped_invoice_payments == 1
    assert rep.imported_count == 15
    assert rep.error_count == 0

    inv_oct = cc.get_invoice(card.id, 2026, 10)
    inv_nov = cc.get_invoice(card.id, 2026, 11)

    assert len(inv_oct.items) == 0
    assert inv_oct.total == 0.00

    assert len(inv_nov.items) == 15
    assert inv_nov.total == 582.74
    assert all("pagamento recebido" not in o.descricao.lower() for o in inv_nov.items)

    # Testa importação de estorno negativo legítimo na mesma fatura
    refund_csv = (
        "date,title,amount\n"
        '2026-10-21,Compra Teste,"100,00"\n'
        '2026-10-22,Estorno de compra,"- 22,74"\n'
    )
    refund_file = tmp_path / "Nubank_2026-11-09_refund.csv"
    refund_file.write_text(refund_csv, encoding="utf-8")
    rep_ref = orchestrator.import_file(
        refund_file,
        profile_id="1",
        target_card_id=card.id,
    )
    assert rep_ref.imported_count == 2
    inv_nov_after = cc.get_invoice(card.id, 2026, 11)
    assert inv_nov_after.total == round(582.74 + 100.00 - 22.74, 2)

    # 1. Tentar reimportar o mesmo CSV sem formatar deve detectar como duplicado
    rep_dup = orchestrator.import_file(
        csv_file,
        profile_id="1",
        target_card_id=card.id,
    )
    assert rep_dup.skipped_duplicates == 15
    assert rep_dup.imported_count == 0

    # 2. Formata o cartão de crédito (limpa despesas e impressões digitais)
    wiped = cc.wipe_card_expenses(card.id)
    assert wiped > 0

    # 3. Agora a reimportação do mesmo arquivo DEVE FUNCIONAR do zero (sem falsos duplicados)
    rep_reimport = orchestrator.import_file(
        csv_file,
        profile_id="1",
        target_card_id=card.id,
    )
    assert rep_reimport.imported_count == 15
    assert rep_reimport.skipped_duplicates == 0

    # 4. Valida botão de Refresh e atalho na MainWindow
    _qapp = _get_qapp()
    backup = BackupSchedulerService(db, settings)
    mw = MainWindow(db, backup)
    assert hasattr(mw, "btn_refresh")
    assert not mw.btn_refresh.icon().isNull() or mw.btn_refresh.text() == "🔄"
    mw.refresh_current_view()
    mw.close()

    db.close()





